# Builds the fictional Word sample offered for download and used in the part 4 clip.
import sys
from docx import Document
from docx.shared import Cm

doc = Document()
doc.add_heading("Atelier Couture Ndèye — Tarifs 2026", level=1)
p = doc.add_paragraph("Voici nos tarifs pour la rentrée. ")
p.add_run("Paiement à 30 jours").bold = True
p.add_run(", par Wave ou en espèces à l'atelier.")
doc.add_heading("Nos prestations", level=2)
for item in ["Uniformes scolaires sur mesure", "Blouses de fête", "Retouches et ourlets"]:
    doc.add_paragraph(item, style="List Bullet")
rows = [("Prestation", "Prix unitaire (F CFA)", "Délai"),
        ("Uniforme (chemise + pantalon)", "7 500", "10 jours"),
        ("Blouse de fête", "5 000", "7 jours"),
        ("Retouche", "1 500", "2 jours")]
t = doc.add_table(rows=len(rows), cols=3)
t.style = "Table Grid"
for r, row in enumerate(rows):
    for c, text in enumerate(row):
        cell = t.cell(r, c)
        cell.text = text
        if r == 0:
            cell.paragraphs[0].runs[0].bold = True
for row in t.rows:
    for cell, w in zip(row.cells, (7, 5, 3)):
        cell.width = Cm(w)
doc.add_paragraph("Document fictif créé pour le tutoriel de l'éditeur Smartacus.")
doc.save(sys.argv[1])
